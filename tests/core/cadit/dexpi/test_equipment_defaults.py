"""Class defaults, nozzle placement and the equipment definition list.

The DEXPI items here are built by hand rather than parsed from a file: the defaults have to answer
for any item the readers can produce, and a synthetic item is the only way to pin a specific shape
(an untagged nozzle, an instrument connection, a chamber) without waiting on a fixture.

Three properties matter more than any individual number. Placement is **deterministic**, or every
re-import diff is noise. No two ports of one equipment share a **position**, or two runs route as
one. Every generated document passes ``validate_equipment_doc``, or the failure surfaces deep
inside the compiler instead of here.
"""

from __future__ import annotations

import json

import pytest

from ada.cadit.dexpi import (
    attributes,
    class_table,
    equipment_defaults,
    equipment_list,
    nozzle_placers,
)
from ada.cadit.dexpi.model import (
    DexpiAttribute,
    DexpiDocument,
    DexpiItem,
    DexpiNode,
    classify,
)
from ada.cadit.dexpi.nozzle_placers import NozzleSpec, place_nozzles
from ada.core.catalog_docs import validate_equipment_doc


# ---------------------------------------------------------------------------
# Synthetic DEXPI items
# ---------------------------------------------------------------------------
def _item(item_id: str, class_name: str, tag: str | None = None, **extra) -> DexpiItem:
    attrs = [DexpiAttribute(name=attributes.TAG_NAME, value=tag)] if tag else []
    return DexpiItem(id=item_id, class_name=class_name, kind=classify(class_name), attributes=attrs, **extra)


def _nozzle(item_id: str, tag: str | None, flow: str | None, class_name: str = "Nozzle", dn: str | None = "100"):
    attrs = []
    if tag:
        attrs.append(DexpiAttribute(name=attributes.SUB_TAG_NAME, value=tag))
    if dn:
        attrs.append(DexpiAttribute(name=attributes.NOMINAL_DIAMETER_NUMERICAL_VALUE_REPRESENTATION, value=dn))
        attrs.append(DexpiAttribute(name=attributes.NOMINAL_DIAMETER_TYPE_REPRESENTATION, value="DN"))
    node = DexpiNode(id=f"{item_id}-node1", ordinal=1, owner_id=item_id, flow=flow)
    return DexpiItem(
        id=item_id,
        class_name=class_name,
        kind=classify(class_name),
        attributes=attrs,
        nodes=[DexpiNode(id=f"{item_id}-node0", ordinal=0, owner_id=item_id, is_anchor=True), node],
    )


def _tiny_doc() -> DexpiDocument:
    """A tank with two nozzles and a level instrument, plus a pump with suction and discharge."""
    doc = DexpiDocument()
    tank = doc.add(_item("Tank1", "Tank", "T-100"))
    doc.add(_nozzle("Tank1-N1", "N1", "in"), tank.id)
    doc.add(_nozzle("Tank1-N2", "N2", "out"), tank.id)
    doc.add(_nozzle("Tank1-N3", "LT1", None, class_name="InstrumentNozzle", dn=None), tank.id)

    pump = doc.add(_item("Pump1", "CentrifugalPump", "P-100"))
    doc.add(_nozzle("Pump1-N1", "S", "in", dn="150"), pump.id)
    doc.add(_nozzle("Pump1-N2", "D", "out", dn="100"), pump.id)
    return doc


# ---------------------------------------------------------------------------
# Class defaults
# ---------------------------------------------------------------------------
def test_meta_disclaims_dexpi_authorship():
    meta = equipment_defaults.meta()
    assert "adapy" in meta["authority"]
    assert "NOT part of the DEXPI specification" in meta["authority"]
    assert "CURATED" in meta["note"]


def test_a_curated_class_answers_for_itself():
    entry = equipment_defaults.resolve_defaults("Tank")
    assert entry["class"] == "Tank"
    assert entry["source"] == "class"
    assert entry["bbox"] == [4.0, 4.0, 6.0]
    assert entry["ifc_element_class"] == "IfcTank"
    assert entry["nozzles"] == "vessel"


@pytest.mark.parametrize(
    "class_name, inherited",
    [
        ("CentrifugalPump", "Pump"),  # one hop
        ("ReciprocatingPump", "Pump"),
        ("PressureVessel", "PressureVessel"),  # curated in its own right
        ("SpiralHeatExchanger", "HeatExchanger"),
        ("Flare", "WasteGasEmitter"),
        ("Conveyor", "StationaryTransportSystem"),
        ("AxialFan", "Fan"),
        ("Boiler", "Heater"),  # two hops: Boiler -> Heater
        ("GasTurbine", "Turbine"),
        ("PackagingSystem", "ProcessEquipment"),  # nothing closer; the catch-all answers
    ],
)
def test_resolution_walks_the_supertype_chain(class_name, inherited):
    assert class_table.get(class_name) is not None, "the class table must know the class under test"
    assert equipment_defaults.curated_classes().count(class_name) == (1 if class_name == inherited else 0)

    entry = equipment_defaults.resolve_defaults(class_name)
    assert entry["class"] == inherited
    assert entry["source"] == ("class" if class_name == inherited else "supertype")


def test_a_few_dozen_entries_cover_every_process_class():
    """The point of walking the DAG: ~29 curated entries answer for all 142 process classes."""
    process = [
        name for name in class_table.class_names() if class_table.get(name)["package"] == "Plant/ProcessEquipment"
    ]
    assert len(process) == 142
    assert len(equipment_defaults.curated_classes()) < 40

    for name in process:
        entry = equipment_defaults.resolve_defaults(name)
        assert all(length > 0 for length in entry["bbox"]), name
        assert entry["ifc_element_class"].startswith("Ifc"), name
        assert entry["nozzles"] in nozzle_placers.PLACERS, name
        assert entry["mass"] > 0, name


def test_chamber_and_nozzle_are_not_process_equipment():
    """PR 1's correction, asserted: neither derives from ProcessEquipment, so neither may be
    classified as equipment. Chamber is curated on its own; Nozzle is never an equipment at all."""
    assert not class_table.is_a("Chamber", "ProcessEquipment")
    assert not class_table.is_a("Nozzle", "ProcessEquipment")
    assert class_table.get("Equipment") is None  # only a Proteus element tag, never a DEXPI class

    assert equipment_defaults.resolve_defaults("Chamber")["source"] == "class"
    assert equipment_defaults.resolve_defaults("Nozzle")["source"] == "fallback"


def test_an_unknown_vendor_class_falls_back_and_says_so():
    entry = equipment_defaults.resolve_defaults("SomeVendorExtensionClass")
    assert entry["class"] is None
    assert entry["source"] == "fallback"
    assert entry["bbox"] == equipment_defaults.resolve_defaults("ProcessEquipment")["bbox"]


def test_mass_follows_an_overridden_envelope():
    """Mass is an envelope density, so a resized box carries a proportionate mass instead of the
    class default's."""
    default = equipment_defaults.build_default_doc("Tank")
    doubled = equipment_defaults.build_default_doc("Tank", bbox=[8.0, 4.0, 6.0])
    assert doubled["mass"] == pytest.approx(default["mass"] * 2.0)


# ---------------------------------------------------------------------------
# Nozzle placement
# ---------------------------------------------------------------------------
def _specs(count: int, direction: str = "INOUT", category: str = "process", prefix: str = "N") -> list[NozzleSpec]:
    return [
        NozzleSpec(
            id=f"{prefix}{i}", name=f"{prefix.lower()}{i}", tag=f"{prefix}{i}", category=category, direction=direction
        )
        for i in range(1, count + 1)
    ]


@pytest.mark.parametrize("strategy", sorted(nozzle_placers.PLACERS))
@pytest.mark.parametrize("count", [1, 2, 3, 7])
def test_placers_are_deterministic_and_never_stack_ports(strategy, count):
    bbox = [4.0, 2.5, 6.0]
    mixed = [
        *_specs(count, "IN", prefix="I"),
        *_specs(count, "OUT", prefix="O"),
        *_specs(count, "INOUT", prefix="U"),
        *_specs(count, "INOUT", "signal", prefix="S"),
        *_specs(count, "INOUT", "electrical", prefix="E"),
    ]
    # Same input twice -> byte-identical output; a reversed input is the same set, so it must place
    # identically too (the placer sorts by ID rather than trusting the caller's order).
    first = place_nozzles(bbox, mixed, strategy)
    assert first == place_nozzles(bbox, mixed, strategy)
    assert first == place_nozzles(bbox, list(reversed(mixed)), strategy)

    positions = [tuple(port["position"]) for port in first]
    assert len(set(positions)) == len(positions)
    assert len({port["name"] for port in first}) == len(first)


def test_ports_sit_on_the_envelope():
    bbox = [4.0, 2.5, 6.0]
    for strategy in sorted(nozzle_placers.PLACERS):
        for port in place_nozzles(bbox, _specs(5, "IN", prefix="I") + _specs(4, "OUT", prefix="O"), strategy):
            x, y, z = port["position"]
            assert -bbox[0] / 2 - 1e-9 <= x <= bbox[0] / 2 + 1e-9
            assert -bbox[1] / 2 - 1e-9 <= y <= bbox[1] / 2 + 1e-9
            assert 0.0 <= z <= bbox[2] + 1e-9
            assert pytest.approx(1.0) == sum(v * v for v in port["direction_vector"])


def test_vessel_feeds_the_top_and_drains_the_side():
    ports = {
        port["name"]: port
        for port in place_nozzles(
            [4.0, 4.0, 6.0],
            [
                NozzleSpec(id="N1", name="feed", direction="IN"),
                NozzleSpec(id="N2", name="bottoms", direction="OUT"),
                NozzleSpec(id="N3", name="side"),
            ],
            "vessel",
        )
    }
    assert ports["feed"]["position"] == [0.0, 0.0, 6.0]
    assert ports["feed"]["direction_vector"] == [0.0, 0.0, 1.0]
    assert ports["bottoms"]["position"][0] == -2.0
    assert ports["bottoms"]["position"][2] < 1.0
    assert ports["side"]["position"][2] == pytest.approx(3.0)


def test_pump_mirrors_the_create_pump_archetype():
    """Suction on -X at mid height, discharge out of the top -- the same geometry
    ``ada.topo_model.equipment.create_pump`` builds, so an imported pump and a hand-built one
    present the same nozzles to the router."""
    ports = {
        port["name"]: port
        for port in place_nozzles(
            [1.5, 0.8, 1.0],
            [
                NozzleSpec(id="N1", name="suction", direction="IN"),
                NozzleSpec(id="N2", name="discharge", direction="OUT"),
            ],
            "pump",
        )
    }
    assert ports["suction"]["position"] == [-0.75, 0.0, 0.5]
    assert ports["suction"]["direction_vector"] == [-1.0, 0.0, 0.0]
    assert ports["discharge"]["position"] == [0.0, 0.0, 1.0]
    assert ports["discharge"]["direction_vector"] == [0.0, 0.0, 1.0]


def test_exchanger_pairs_nozzles_on_the_two_ends():
    ports = place_nozzles([6.0, 1.2, 1.2], _specs(4), "exchanger")
    faces = [port["direction_vector"][0] for port in ports]
    assert faces == [-1.0, 1.0, -1.0, 1.0]


def test_duplicate_tags_are_uniquified_not_dropped():
    ports = place_nozzles([2.0, 2.0, 2.0], [NozzleSpec(id=f"N{i}", name="n", tag="N") for i in (1, 2, 3)], "generic")
    assert [port["name"] for port in ports] == ["n", "n-2", "n-3"]


def test_unknown_strategy_is_loud():
    with pytest.raises(ValueError, match="unknown nozzle-layout strategy"):
        place_nozzles([1.0, 1.0, 1.0], _specs(1), "helicoidal")


# ---------------------------------------------------------------------------
# Categories -- a wrong one silently drops a whole system at compile time
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "class_name, expected",
    [
        ("Nozzle", "process"),
        ("ProcessNozzle", "process"),
        ("AccessNozzle", "process"),
        ("InstrumentNozzle", "signal"),
        ("ProcessInstrumentationFunction", "signal"),
        ("SignalOffPageConnector", "signal"),
        ("SignalLineFunction", "signal"),
        ("ActuatingElectricalSystem", "electrical"),
        ("ActuatingElectricalFunction", "electrical"),
        ("SomeVendorExtensionClass", "process"),
    ],
)
def test_category_from_class(class_name, expected):
    assert nozzle_placers.category_for(class_name) == expected


def test_a_plain_nozzle_is_not_electrical_despite_its_supertypes():
    """The trap: the spec makes ``Nozzle`` an ``ActuatingElectricalLocation`` *and* a
    ``SensingLocation``, so a naive supertype test types every process nozzle as electrical -- and
    ``System.connect`` then refuses the piping run and the compiler drops the system with a
    warning."""
    assert class_table.is_a("Nozzle", "ActuatingElectricalLocation")
    assert nozzle_placers.category_for("Nozzle") == "process"


@pytest.mark.parametrize(
    "node_type, expected",
    [("process", "process"), ("signal", "signal"), ("electrical", "electrical"), ("", "process"), (None, "process")],
)
def test_category_from_node_type(node_type, expected):
    node = DexpiNode(id="n1", ordinal=1, node_type=node_type)
    assert nozzle_placers.nozzle_from_node(node).category == expected


@pytest.mark.parametrize(
    "flow, expected", [("in", "IN"), ("out", "OUT"), ("OUT", "OUT"), (None, "INOUT"), ("unknown", "INOUT")]
)
def test_direction_from_flow(flow, expected):
    assert nozzle_placers.direction_for(flow) == expected


def test_nozzle_identity_rides_onto_the_port():
    doc = _tiny_doc()
    specs = equipment_list.nozzle_specs_for(doc, doc.items["Pump1"])
    ports = {port["tag"]: port for port in place_nozzles([1.5, 0.8, 1.0], specs, "pump")}
    assert ports["S"]["nominal_diameter"] == pytest.approx(0.15)
    assert ports["D"]["nominal_diameter"] == pytest.approx(0.1)
    assert ports["S"]["direction"] == "IN"
    assert ports["D"]["direction"] == "OUT"


# ---------------------------------------------------------------------------
# The definition list
# ---------------------------------------------------------------------------
def test_merge_generates_a_valid_doc_per_equipment():
    definitions = equipment_list.merge_definitions(_tiny_doc())
    assert sorted(definitions) == ["p-100", "t-100"]

    for slug, doc in definitions.items():
        assert validate_equipment_doc(doc) == doc, slug

    tank = definitions["t-100"]
    assert tank["dexpi_class"] == "Tank"
    assert tank["dexpi_id"] == "Tank1"
    assert tank["tag"] == "T-100"
    assert tank["ifc_element_class"] == "IfcTank"
    assert [port["category"] for port in tank["ports"]] == ["process", "process", "signal"]

    pump = definitions["p-100"]
    assert pump["dexpi_class"] == "CentrifugalPump"  # provenance is the item's own class...
    assert pump["ifc_element_class"] == "IfcPump"  # ...while the defaults came from Pump


def test_chamber_nozzles_are_folded_into_their_owner():
    doc = DexpiDocument()
    separator = doc.add(_item("V1", "Separator", "V-201"))
    doc.add(_nozzle("V1-N1", "N1", "in"), separator.id)
    boot = doc.add(_item("V1-C1", "Chamber"), separator.id)
    doc.add(_nozzle("V1-C1-N1", "N9", "out"), boot.id)

    definitions = equipment_list.merge_definitions(doc)
    assert list(definitions) == ["v-201"]  # the chamber is not an equipment of its own
    assert sorted(port["tag"] for port in definitions["v-201"]["ports"]) == ["N1", "N9"]


def test_per_tag_beats_per_class_beats_default():
    doc = _tiny_doc()
    default = equipment_list.merge_definitions(doc)["t-100"]["bbox"]
    assert default == {"lx": 4.0, "ly": 4.0, "lz": 6.0}

    by_class = equipment_list.merge_definitions(doc, {"Tank": {"bbox": {"lx": 3.0, "ly": 3.0, "lz": 9.0}}})
    assert by_class["t-100"]["bbox"] == {"lx": 3.0, "ly": 3.0, "lz": 9.0}

    both = equipment_list.merge_definitions(
        doc,
        {
            "Tank": {"bbox": {"lx": 3.0, "ly": 3.0, "lz": 9.0}, "mass": 1.0},
            "T-100": {"bbox": {"lx": 5.0, "ly": 5.0, "lz": 12.0}},
        },
    )
    assert both["t-100"]["bbox"] == {"lx": 5.0, "ly": 5.0, "lz": 12.0}
    assert both["t-100"]["mass"] == 1.0  # the per-class override still supplies what the tag omits
    # The other tank-less item is untouched by either override.
    assert both["p-100"]["bbox"] == {"lx": 1.5, "ly": 0.8, "lz": 1.0}


def test_ports_are_generated_against_the_overridden_envelope():
    """An override that only resizes the box must not leave the nozzles floating off it."""
    doc = _tiny_doc()
    definitions = equipment_list.merge_definitions(doc, {"T-100": {"bbox": {"lx": 10.0, "ly": 10.0, "lz": 20.0}}})
    feed = next(port for port in definitions["t-100"]["ports"] if port["direction"] == "IN")
    assert feed["position"] == [0.0, 0.0, 20.0]


def test_an_override_listing_ports_replaces_them():
    doc = _tiny_doc()
    override = {"ports": [{"name": "only", "position": [0.0, 0.0, 1.0], "direction_vector": [0.0, 0.0, 1.0]}]}
    definitions = equipment_list.merge_definitions(doc, {"T-100": override})
    assert [port["name"] for port in definitions["t-100"]["ports"]] == ["only"]


def test_the_resolver_is_just_a_dict_get():
    resolver = equipment_list.dexpi_equipment_resolver(_tiny_doc())
    assert resolver("t-100")["ifc_element_class"] == "IfcTank"
    assert resolver("no-such-slug") is None


def test_json_and_xlsx_loaders_round_trip_the_same_dict(tmp_path):
    definitions = equipment_list.merge_definitions(_tiny_doc())

    as_json = tmp_path / "definitions.json"
    as_xlsx = tmp_path / "definitions.xlsx"
    equipment_list.write_equipment_definitions(definitions, as_json)
    equipment_list.write_equipment_definitions(definitions, as_xlsx)

    from_json = equipment_list.load_equipment_definitions(as_json)
    from_xlsx = equipment_list.load_equipment_definitions(as_xlsx)
    assert from_json == definitions
    assert from_xlsx == from_json

    # Two sheets joined on SLUG, so the nozzles stay hand-editable rows.
    from openpyxl import load_workbook

    workbook = load_workbook(as_xlsx)
    assert set(workbook.sheetnames) == {"EquipmentTypes", "Nozzles"}
    assert workbook["Nozzles"].cell(row=1, column=1).value == "SLUG"
    assert workbook["Nozzles"].max_row == 1 + sum(len(doc["ports"]) for doc in definitions.values())


def test_a_hand_written_json_list_loads_and_validates(tmp_path):
    path = tmp_path / "hand.json"
    path.write_text(json.dumps({"V-201": {"bbox": {"lx": 6.0, "ly": 3.0, "lz": 3.0}}}), encoding="utf-8")
    loaded = equipment_list.load_equipment_definitions(path)
    assert loaded["V-201"]["bbox"] == {"lx": 6.0, "ly": 3.0, "lz": 3.0}
    assert loaded["V-201"]["mass"] == 1000.0  # the catalog default, untouched


def test_a_generated_doc_builds_a_real_equipment():
    """The other end of the contract: the documents this module emits are what
    ``build_equipment_from_catalog`` consumes, port identity included."""
    from ada.topo_model.equipment import build_equipment_from_catalog

    doc = equipment_list.merge_definitions(_tiny_doc())["p-100"]
    eq = build_equipment_from_catalog("P-100", (0.0, 0.0, 0.0), doc)
    assert eq.ifc_element_class == "IfcPump"
    assert eq.mass == pytest.approx(doc["mass"])

    suction = eq.get_port("s")
    assert suction.tag == "S"
    assert suction.nominal_diameter == pytest.approx(0.15)
    assert suction.direction.value == "IN"
    assert suction.category == "process"


def test_a_bad_definition_list_is_rejected(tmp_path):
    path = tmp_path / "bad.json"
    path.write_text(json.dumps({"V-201": {"ports": [{"name": "a"}, {"name": "a"}]}}), encoding="utf-8")
    with pytest.raises(ValueError, match="duplicate port names"):
        equipment_list.load_equipment_definitions(path)

    with pytest.raises(ValueError, match=r"must be \.json or \.xlsx"):
        equipment_list.load_equipment_definitions(tmp_path / "definitions.csv")
